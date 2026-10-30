import sys
from pptx import Presentation

p2 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289566441_6b39c740.pptx'
prs = Presentation(p2)

for i, s in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {i} ====================")
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    runs_info = []
                    for r in p.runs:
                        color_str = ""
                        try:
                            if r.font.color and r.font.color.rgb:
                                color_str = f"#{r.font.color.rgb}"
                        except:
                            pass
                        runs_info.append(f"['{r.text}', bold={r.font.bold}, size={r.font.size.pt if r.font.size else None}, color={color_str}]")
                    sys.stdout.buffer.write((f"  SHAPE {sh.shape_id} ('{sh.name}'): {txt}\n    RUNS: {' + '.join(runs_info)}\n").encode('utf-8'))
        elif sh.shape_type == 13: # PICTURE
            print(f"  PICTURE {sh.shape_id} ('{sh.name}'): pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), size=({sh.width.inches:.2f}, {sh.height.inches:.2f})")
