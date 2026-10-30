import os
from pptx import Presentation

p1 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289906735_311b416e.pptx'
prs = Presentation(p1)

print("Slide count:", len(prs.slides))
# Check if we can delete slides from the end using internal xml
def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[index])

# Test deleting slide 8 (icons)
print("Before delete:", len(prs.slides))
delete_slide(prs, 8)
print("After delete:", len(prs.slides))
delete_slide(prs, 7)
print("After second delete:", len(prs.slides))

out = r'C:\SPARC\docs\defesa\test_deleted.pptx'
prs.save(out)
print("Saved test_deleted.pptx successfully!")
