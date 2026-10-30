from pptx import Presentation

p2 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289566441_6b39c740.pptx'
prs2 = Presentation(p2)
for i in range(3):
    s = prs2.slides[i]
    print(f'=== PPT2 Slide {i+1}: layout="{s.slide_layout.name}", shapes={len(s.shapes)} ===')
    for sh in s.shapes:
        txt = sh.text_frame.text.replace('\n', ' // ') if sh.has_text_frame else ''
        txt_clean = txt[:60].encode('ascii', errors='replace').decode('ascii')
        print(f'  name="{sh.name}", type={sh.shape_type}, left={sh.left.inches:.2f}, top={sh.top.inches:.2f}, w={sh.width.inches:.2f}, h={sh.height.inches:.2f}, txt="{txt_clean}"')
