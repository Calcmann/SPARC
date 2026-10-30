from pptx import Presentation

p = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
prs = Presentation(p)

for idx in range(len(prs.slides)):
    s = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for sh in s.shapes:
        if sh.has_text_frame:
            print(f"Shape ID={sh.shape_id}, Name='{sh.name}', Pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), Size=({sh.width.inches:.2f}, {sh.height.inches:.2f})")
            for p_idx, p_para in enumerate(sh.text_frame.paragraphs):
                runs_t = [r.text for r in p_para.runs]
                print(f"   P{p_idx}: text='{p_para.text}' | runs={runs_t}")
