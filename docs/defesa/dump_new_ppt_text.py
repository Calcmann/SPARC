import sys
from pptx import Presentation

p = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
prs = Presentation(p)

for idx, s in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {idx} ====================")
    for sh in s.shapes:
        if sh.has_text_frame:
            for p_para in sh.text_frame.paragraphs:
                txt = p_para.text.strip()
                if txt:
                    sys.stdout.buffer.write((f"  [{sh.name}] {txt}\n").encode('utf-8'))
        elif sh.shape_type == 6: # GROUP
            for sub_sh in sh.shapes:
                if sub_sh.has_text_frame:
                    for p_para in sub_sh.text_frame.paragraphs:
                        txt = p_para.text.strip()
                        if txt:
                            sys.stdout.buffer.write((f"  [GROUP {sh.name} -> {sub_sh.name}] {txt}\n").encode('utf-8'))
