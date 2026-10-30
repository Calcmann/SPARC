from pptx import Presentation

p1 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289906735_311b416e.pptx'
prs1 = Presentation(p1)
for i in range(7):
    s = prs1.slides[i]
    print(f'=== Slide {i+1}: layout="{s.slide_layout.name}", shapes={len(s.shapes)} ===')
    for sh in s.shapes:
        txt = sh.text_frame.text.replace('\n', ' // ') if sh.has_text_frame else ''
        txt_clean = txt[:60].encode('ascii', errors='replace').decode('ascii')
        print(f'  name="{sh.name}", type={sh.shape_type}, left={sh.left.inches:.2f}, top={sh.top.inches:.2f}, w={sh.width.inches:.2f}, h={sh.height.inches:.2f}, txt="{txt_clean}"')
