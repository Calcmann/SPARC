from pptx import Presentation

p1 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289906735_311b416e.pptx'
prs = Presentation(p1)

for i in range(len(prs.slides)):
    s = prs.slides[i]
    print(f"Slide {i+1}:")
    for sh in s.shapes:
        print(f"  Shape id={sh.shape_id}, name='{sh.name}', type={sh.shape_type}, left={sh.left.inches:.2f}, top={sh.top.inches:.2f}, w={sh.width.inches:.2f}, h={sh.height.inches:.2f}")
