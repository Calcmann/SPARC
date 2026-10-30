import sys, os
from pptx import Presentation

p = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
prs = Presentation(p)

print(f"=== NEW PPTX: {os.path.basename(p)} ===")
print(f"Total slides: {len(prs.slides)}, Width: {prs.slide_width.inches:.2f}, Height: {prs.slide_height.inches:.2f}")

for idx, s in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {idx} (layout: '{s.slide_layout.name}') ====================")
    for sh in s.shapes:
        info = f"  Shape {sh.shape_id} ('{sh.name}'): type={sh.shape_type}, pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), size=({sh.width.inches:.2f}, {sh.height.inches:.2f})"
        if sh.has_text_frame:
            txt = sh.text_frame.text.replace('\n', ' // ').strip()
            info += f", text='{txt[:80]}'"
        sys.stdout.buffer.write((info + '\n').encode('utf-8'))
