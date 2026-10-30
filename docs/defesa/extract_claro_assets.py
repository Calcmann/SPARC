import os
from pptx import Presentation

p1 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289906735_311b416e.pptx'
prs = Presentation(p1)

extract_dir = r'C:\SPARC\docs\defesa\claro_template_assets'
os.makedirs(extract_dir, exist_ok=True)

img_count = 0
for s_idx, slide in enumerate(prs.slides, 1):
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.shape_type == 13: # PICTURE
            image = shape.image
            filename = f"slide_{s_idx}_shape_{shape.shape_id}_{shape.name}.{image.ext}"
            filename = filename.replace(' ', '_').replace('?', '_')
            target = os.path.join(extract_dir, filename)
            with open(target, 'wb') as f:
                f.write(image.blob)
            print(f"Slide {s_idx}: Extracted {filename} ({len(image.blob)} bytes, {image.size})")
            img_count += 1

print(f"Total extracted: {img_count}")
