import os
from pptx import Presentation

p1 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289906735_311b416e.pptx'
prs = Presentation(p1)

print(f"Total slides: {len(prs.slides)}")
for i, s in enumerate(prs.slides, 1):
    print(f"\n--- SLIDE {i} ---")
    for sh in s.shapes:
        txt = sh.text_frame.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        print(f"  Shape {sh.shape_id} ('{sh.name}'): type={sh.shape_type}, pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), size=({sh.width.inches:.2f}, {sh.height.inches:.2f}), txt='{txt[:50]}'")
