import sys
from pptx import Presentation

p = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
prs = Presentation(p)

for idx in range(2, len(prs.slides)):
    s = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for sh in s.shapes:
        txt = sh.text_frame.text.replace('\n', ' // ').strip() if sh.has_text_frame else ''
        print(f"Shape {sh.shape_id} ('{sh.name}'): pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), size=({sh.width.inches:.2f}, {sh.height.inches:.2f}), type={sh.shape_type}")
        if txt:
            print(f"   TXT: {txt}")
        if sh.shape_type == 6: # GROUP
            for sub in sh.shapes:
                sub_txt = sub.text_frame.text.replace('\n', ' // ').strip() if sub.has_text_frame else ''
                print(f"   SUB {sub.shape_id} ('{sub.name}'): pos=({sub.left.inches:.2f}, {sub.top.inches:.2f}), size=({sub.width.inches:.2f}, {sub.height.inches:.2f}), type={sub.shape_type}, txt='{sub_txt}'")
