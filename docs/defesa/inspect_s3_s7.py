from pptx import Presentation

p = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
prs = Presentation(p)

for idx in [2, 3, 4, 5, 6]:
    s = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for sh in s.shapes:
        if sh.has_text_frame:
            print(f"Shape ID={sh.shape_id}, Name='{sh.name}', Pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), Size=({sh.width.inches:.2f}, {sh.height.inches:.2f})")
            for p_idx, p_para in enumerate(sh.text_frame.paragraphs):
                print(f"   P{p_idx}: text='{p_para.text}'")
        elif sh.shape_type == 6: # GROUP
            print(f"GROUP Shape ID={sh.shape_id}, Name='{sh.name}'")
            for sub in sh.shapes:
                if sub.has_text_frame:
                    print(f"   SUB {sub.shape_id} ('{sub.name}'): txt='{sub.text_frame.text}'")
