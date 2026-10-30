import os
from pptx import Presentation

p2 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289566441_6b39c740.pptx'
prs2 = Presentation(p2)

out_dir = r'C:\SPARC\docs\defesa\user_ppt_assets'
os.makedirs(out_dir, exist_ok=True)

for i, s in enumerate(prs2.slides, 1):
    for sh in s.shapes:
        if sh.shape_type == 13: # PICTURE
            img = sh.image
            fn = f"slide_{i}_{sh.shape_id}_{sh.name}.{img.ext}"
            fn = fn.replace(' ', '_').replace('?', '_')
            path = os.path.join(out_dir, fn)
            with open(path, 'wb') as f:
                f.write(img.blob)
            print(f"Extracted Slide {i}: {fn} ({len(img.blob)} bytes, {img.size})")
