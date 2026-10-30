import sys
from pptx import Presentation

p = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
prs = Presentation(p)

for i in [0, 1]:
    s = prs.slides[i]
    print(f"\n==================== SLIDE {i+1} ====================")
    for sh in s.shapes:
        txt = sh.text_frame.text.replace('\n', ' // ') if sh.has_text_frame else ''
        print(f"  Shape {sh.shape_id} ('{sh.name}'): pos=({sh.left.inches:.2f}, {sh.top.inches:.2f}), size=({sh.width.inches:.2f}, {sh.height.inches:.2f}), txt='{txt}'")
