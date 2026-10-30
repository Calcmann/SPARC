from pptx import Presentation

p1 = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793289906735_311b416e.pptx'
prs = Presentation(p1)

for idx, layout in enumerate(prs.slide_layouts):
    print(f"Layout {idx}: '{layout.name}', shapes in layout={len(layout.shapes)}")
    for sh in layout.shapes:
        print(f"   shape: name='{sh.name}', type={sh.shape_type}")
