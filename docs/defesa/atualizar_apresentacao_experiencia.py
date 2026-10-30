import os, sys, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC_PPT = r'C:\Users\User\.gemini\antigravity\brain\4c8109a0-5840-4810-a56c-709d56ac8b9f\.user_uploaded\media_1793307933366_60bf5cd5.pptx'
OUT_PPT = r'C:\SPARC\docs\defesa\SPARC_Apresentacao_Inovacao_Claro.pptx'
DIST_PPT = r'C:\SPARC\dist-beta\SPARC_Apresentacao_Inovacao_Claro.pptx'

IMG_DESKTOP_UI = r'C:\SPARC\docs\defesa\user_ppt_assets\slide_2_36_Imagem_35.png'
IMG_DESKTOP_16_9 = r'C:\SPARC\docs\defesa\assets\sparc_ui_16_9.png'
IMG_MOBILE_PROV = r'C:\SPARC\docs\defesa\user_ppt_assets\slide_3_21_Imagem_20.png'
IMG_MOBILE_TEST = r'C:\SPARC\docs\defesa\user_ppt_assets\slide_3_23_Imagem_22.png'

# Cores oficiais Claro
CLARO_RED = RGBColor(0xDA, 0x29, 0x1C)
CLARO_DARK = RGBColor(0x26, 0x26, 0x26)
CLARO_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLARO_GRAY = RGBColor(0x66, 0x66, 0x66)
CLARO_LIGHT_BG = RGBColor(0xF8, 0xF9, 0xFA)
CLARO_BORDER = RGBColor(0xE0, 0xE0, 0xE0)
CLARO_YELLOW = RGBColor(0xFF, 0xC1, 0x07)
CLARO_ORANGE = RGBColor(0xF5, 0x84, 0x2B)

def set_shape_image(slide, shape, image_path):
    """Substitui a imagem de uma forma de figura sem quebrar os relacionamentos do PPTX."""
    img_part, rId = slide.part.get_or_add_image_part(image_path)
    blip = shape._element.blipFill.blip
    blip.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'] = rId

def clear_text_frame(tf):
    """Limpa todos os parágrafos de um text frame, mantendo apenas um vazio."""
    p_elems = list(tf._txBody.p_lst)
    for p in p_elems[1:]:
        tf._txBody.remove(p)
    if tf.paragraphs:
        tf.paragraphs[0].text = ""

def remove_shape(slide, shape):
    """Remove uma forma do slide."""
    sp = shape._element
    sp.getparent().remove(sp)

def main():
    print(f"Copiando template oficial Claro...")
    shutil.copy(SRC_PPT, OUT_PPT)
    prs = Presentation(OUT_PPT)
    print(f"Apresentação carregada: {len(prs.slides)} slides.")

    # ==========================================
    # SLIDE 1: Capa (Título & Categoria)
    # ==========================================
    print("Atualizando Slide 1 (Capa)...")
    s1 = prs.slides[0]
    sh_title = [sh for sh in s1.shapes if sh.shape_id == 8][0]
    clear_text_frame(sh_title.text_frame)
    p = sh_title.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Título do Projeto: "
    r1.font.name = "AMX Medium"
    r1.font.size = Pt(24)
    r1.font.bold = True
    r1.font.color.rgb = CLARO_RED

    r2 = p.add_run()
    r2.text = "SPARC — Sistema de Provisionamento e Ativação de Roteadores Claro"
    r2.font.name = "AMX Medium"
    r2.font.size = Pt(24)
    r2.font.bold = True
    r2.font.color.rgb = CLARO_WHITE

    sh_cat = [sh for sh in s1.shapes if sh.shape_id == 9][0]
    clear_text_frame(sh_cat.text_frame)
    p = sh_cat.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Categoria: "
    r1.font.name = "AMX Medium"
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = CLARO_RED

    r2 = p.add_run()
    r2.text = "Ideia que melhora a experiência do cliente"
    r2.font.name = "AMX Medium"
    r2.font.size = Pt(22)
    r2.font.color.rgb = CLARO_WHITE

    # ==========================================
    # SLIDE 2: Integrantes do Projeto (MANTER CONTEÚDO)
    # ==========================================
    print("Slide 2 preservado intacto (George Luís Calcmann).")

    # ==========================================
    # SLIDE 3: Objetivo e Justificativa
    # ==========================================
    print("Atualizando Slide 3 (Objetivo e Justificativa)...")
    s3 = prs.slides[2]
    # Shape 4: Texto Principal
    sh_obj = [sh for sh in s3.shapes if sh.shape_id == 4][0]
    clear_text_frame(sh_obj.text_frame)
    
    p1 = sh_obj.text_frame.paragraphs[0]
    p1.space_after = Pt(10)
    r = p1.add_run()
    r.text = "Automatizar integralmente a esteira de detecção, preparação física, gravação de firmware e provisionamento de roteadores corporativos B2B (Cisco, HPE e Fortinet), eliminando processos manuais em linha de comando que levavam até 50 minutos por equipamento."
    r.font.name = "AMX"
    r.font.size = Pt(14)
    r.font.color.rgb = CLARO_DARK

    p2 = sh_obj.text_frame.add_paragraph()
    p2.space_after = Pt(10)
    r = p2.add_run()
    r.text = "Garantir o cumprimento rigoroso da Visita Única (VU), reduzindo o tempo de ativação para menos de 10 minutos na ponta, eliminando erros humanos de digitação de parâmetros do SAIP e expurgando configurações residuais de clientes anteriores que causavam falsos diagnósticos de rede."
    r.font.name = "AMX"
    r.font.size = Pt(14)
    r.font.color.rgb = CLARO_DARK

    p3 = sh_obj.text_frame.add_paragraph()
    r = p3.add_run()
    r.text = "A estratégia central transfere o esforço técnico de configuração para a base Claro (modelo Plug & Play), blindando a operação de campo, reduzindo custos com terceiros e assegurando satisfação e confiabilidade desde o primeiro minuto de ativação."
    r.font.name = "AMX"
    r.font.size = Pt(14)
    r.font.color.rgb = CLARO_DARK

    # 3 Cards Inferiores:
    # Card 1 (Shape 8)
    sh_c1 = [sh for sh in s3.shapes if sh.shape_id == 8][0]
    clear_text_frame(sh_c1.text_frame)
    p = sh_c1.text_frame.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run()
    r.text = "Garantir Visita Única (VU)\n"
    r.font.name = "AMX Black"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = CLARO_RED
    r2 = p.add_run()
    r2.text = "Ativação em <10 min no rack sem retrabalho, gerando laudo digital imediato com QR Code."
    r2.font.name = "AMX"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = CLARO_DARK

    # Card 2 (Shape 20)
    sh_c2 = [sh for sh in s3.shapes if sh.shape_id == 20][0]
    clear_text_frame(sh_c2.text_frame)
    p = sh_c2.text_frame.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run()
    r.text = "Triagem & Plug & Play\n"
    r.font.name = "AMX Black"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = CLARO_RED
    r2 = p.add_run()
    r2.text = "Roteador sai 100% testado da base (portas, memórias, boot) e com firmware gravado."
    r2.font.name = "AMX"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = CLARO_DARK

    # Card 3 (Shape 22)
    sh_c3 = [sh for sh in s3.shapes if sh.shape_id == 22][0]
    clear_text_frame(sh_c3.text_frame)
    p = sh_c3.text_frame.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run()
    r.text = "Custo Zero & Soberania\n"
    r.font.name = "AMX Black"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = CLARO_RED
    r2 = p.add_run()
    r2.text = "100% desenvolvido internamente, sem licenças, sem novos servidores e Offline-First."
    r2.font.name = "AMX"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = CLARO_DARK

    # Imagem à direita (Shape 1026) -> Substituir com screenshot mobile
    sh_pic2 = [sh for sh in s3.shapes if sh.shape_id == 1026][0]
    set_shape_image(s3, sh_pic2, IMG_MOBILE_PROV)

    # ==========================================
    # SLIDE 4: Desafios | Cenário da Problemática
    # ==========================================
    print("Atualizando Slide 4 (Desafios)...")
    s4 = prs.slides[3]
    
    # Lista de desafios SPARC
    desafios = [
        (6, "1. Tempo excessivo e ativação manual (até 50 min/roteador)",
         "Processo artesanal linha por linha em terminal serial, gerando filas de espera e ociosidade da equipe técnica de campo."),
        (22, "2. Erros humanos de digitação e falhas no SAIP",
         "Transcrição manual de IPs de WAN, máscaras e rotas estáticas; erro de 1 único dígito travava a ativação e exigia retrabalho."),
        (24, "3. Visitas perdidas por defeitos físicos não detectados",
         "Roteadores do estoque com portas queimadas ou memórias corrompidas só eram descobertos na frente do cliente B2B."),
        (25, "4. Contaminação por configurações residuais (VLANs antigas)",
         "Equipamentos reutilizados retinham senhas e VLANs de clientes anteriores, gerando falso diagnóstico de 'link down'."),
        (23, "5. Dependência de conectividade externa e custos com terceiros",
         "Ferramentas de mercado (ex: Panteon) exigiam link de internet ativo e servidores centrais, inviáveis em subsolos e obras.")
    ]

    for sh_id, title_txt, desc_txt in desafios:
        sh = [s for s in s4.shapes if s.shape_id == sh_id][0]
        sh.width = Inches(11.2)
        sh.height = Inches(0.80)
        clear_text_frame(sh.text_frame)
        p = sh.text_frame.paragraphs[0]
        r1 = p.add_run()
        r1.text = title_txt + "\n"
        r1.font.name = "AMX Black"
        r1.font.size = Pt(13.5)
        r1.font.bold = True
        r1.font.color.rgb = CLARO_RED

        r2 = p.add_run()
        r2.text = desc_txt
        r2.font.name = "AMX"
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 5: Ações Desenvolvidas
    # ==========================================
    print("Atualizando Slide 5 (Ações Desenvolvidas)...")
    s5 = prs.slides[4]
    grp22 = [sh for sh in s5.shapes if sh.name == 'Agrupar 22'][0]

    acoes = [
        ('CaixaDeTexto 10', "1. Automação SAIP", "Digitalização dos fluxos e leitor inteligente do SAIP (PDF/TXT), eliminando digitação manual de parâmetros."),
        ('CaixaDeTexto 69', "2. TFTP Nativo & Offline", "Servidor gigabit embutido com cache local de imagens; 100% autônomo sem internet."),
        ('CaixaDeTexto 73', "3. Drivers Multi-Marca", "Isolamento por modelo (Cisco IOS/IOS-XE, HPE Comware, Fortinet) com quebra de senha."),
        ('CaixaDeTexto 77', "4. Solução Híbrida", "SPARC Desktop para bancada e SPARC Mobile Android via cabo OTG para técnicos na ponta."),
        ('CaixaDeTexto 81', "5. Triagem e 789 Testes", "Burn-in de portas e memórias na base e qualidade blindada por 789 testes unitários automatizados (0 falhas).")
    ]

    for box_name, title_txt, desc_txt in acoes:
        sh = [sub for sub in grp22.shapes if sub.name == box_name][0]
        clear_text_frame(sh.text_frame)
        p = sh.text_frame.paragraphs[0]
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = title_txt + "\n\n"
        r1.font.name = "AMX Black"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = CLARO_RED

        r2 = p.add_run()
        r2.text = desc_txt
        r2.font.name = "AMX"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 6: A Solução: Esteira Automatizada SPARC
    # ==========================================
    print("Atualizando Slide 6 (A Solução - Esteira)...")
    s6 = prs.slides[5]
    sh_t6 = [sh for sh in s6.shapes if sh.shape_id == 39][0]
    sh_t6.text_frame.paragraphs[0].text = "A Solução: Esteira Automatizada SPARC"
    sh_t6.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_t6.text_frame.paragraphs[0].font.size = Pt(36)
    sh_t6.text_frame.paragraphs[0].font.bold = True

    # Imagem 1 no centro -> Substituir com screenshot Desktop UI
    sh_img1 = [sh for sh in s6.shapes if sh.shape_id == 2][0]
    set_shape_image(s6, sh_img1, IMG_DESKTOP_UI)

    # ==========================================
    # SLIDE 7: SPARC Desktop: Preparação e Triagem na Base Claro
    # ==========================================
    print("Atualizando Slide 7 (SPARC Desktop)...")
    s7 = prs.slides[6]
    sh_t7 = [sh for sh in s7.shapes if sh.shape_id == 39][0]
    sh_t7.text_frame.paragraphs[0].text = "SPARC Desktop: Preparação e Triagem na Base Claro"
    sh_t7.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_t7.text_frame.paragraphs[0].font.size = Pt(32)
    sh_t7.text_frame.paragraphs[0].font.bold = True

    # Imagem 7 -> Mudar para esquerda e colocar UI Widescreen
    sh_img7 = [sh for sh in s7.shapes if sh.shape_id == 8][0]
    sh_img7.left = Inches(0.85)
    sh_img7.top = Inches(1.5)
    sh_img7.width = Inches(6.8)
    sh_img7.height = Inches(5.0)
    set_shape_image(s7, sh_img7, IMG_DESKTOP_16_9)

    # Remover textos antigos do slide 7
    sh_txt_old1 = [sh for sh in s7.shapes if sh.shape_id == 2][0]
    remove_shape(s7, sh_txt_old1)
    
    sh_txt2 = [sh for sh in s7.shapes if sh.shape_id == 3][0]
    sh_txt2.left = Inches(7.85)
    sh_txt2.top = Inches(1.5)
    sh_txt2.width = Inches(4.8)
    sh_txt2.height = Inches(5.0)
    clear_text_frame(sh_txt2.text_frame)

    itens_desktop = [
        ("Detecção Automática Plug & Play", "Identificação instantânea do modelo (Cisco, HPE, Fortinet), portas e versão via console serial."),
        ("Carga de Firmware Gigabit", "Servidor TFTP nativo grava 500 MB em ~3 min (substituindo transferências lentas de 35 min)."),
        ("Higienização Profunda (Clean State)", "Expurgo de configurações residuais, senhas antigas e VLANs de clientes anteriores."),
        ("Testes Fisiológicos de Hardware", "Teste de integridade de memória RAM, integridade da Flash e loopback de portas Ethernet."),
        ("Saída Homologada para Campo", "Roteador sai da base 100% testado, etiquetado e Plug & Play para instalação imediata no rack.")
    ]

    for idx, (titulo, desc) in enumerate(itens_desktop):
        p = sh_txt2.text_frame.paragraphs[0] if idx == 0 else sh_txt2.text_frame.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + titulo + "\n  "
        r1.font.name = "AMX Black"
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = CLARO_RED

        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "AMX"
        r2.font.size = Pt(11)
        r2.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 8: SPARC Mobile Android: Validação e Laudo na Ponta
    # ==========================================
    print("Atualizando Slide 8 (SPARC Mobile Android)...")
    s8 = prs.slides[7]
    sh_t8 = [sh for sh in s8.shapes if sh.shape_id == 39][0]
    sh_t8.text_frame.paragraphs[0].text = "SPARC Mobile Android: Validação e Laudo na Ponta"
    sh_t8.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_t8.text_frame.paragraphs[0].font.size = Pt(32)
    sh_t8.text_frame.paragraphs[0].font.bold = True

    # Imagem 7 (left phone) -> Provisioning screen
    sh_m1 = [sh for sh in s8.shapes if sh.shape_id == 8][0]
    sh_m1.left = Inches(0.85)
    sh_m1.top = Inches(1.4)
    sh_m1.width = Inches(2.7)
    sh_m1.height = Inches(5.2)
    set_shape_image(s8, sh_m1, IMG_MOBILE_PROV)

    # Imagem 8 (middle phone) -> Speedtest & Report screen
    sh_m2 = [sh for sh in s8.shapes if sh.shape_id == 9][0]
    sh_m2.left = Inches(3.85)
    sh_m2.top = Inches(1.4)
    sh_m2.width = Inches(2.7)
    sh_m2.height = Inches(5.2)
    set_shape_image(s8, sh_m2, IMG_MOBILE_TEST)

    # Limpar/remover formas de texto antigas e imagem 12
    sh_m3 = [sh for sh in s8.shapes if sh.shape_id == 13][0]
    remove_shape(s8, sh_m3)
    sh_old_t1 = [sh for sh in s8.shapes if sh.shape_id == 3][0]
    remove_shape(s8, sh_old_t1)
    sh_old_t2 = [sh for sh in s8.shapes if sh.shape_id == 11][0]
    remove_shape(s8, sh_old_t2)

    sh_card_m = [sh for sh in s8.shapes if sh.shape_id == 15][0]
    sh_card_m.left = Inches(6.9)
    sh_card_m.top = Inches(1.4)
    sh_card_m.width = Inches(5.8)
    sh_card_m.height = Inches(5.2)
    clear_text_frame(sh_card_m.text_frame)

    itens_mobile = [
        ("Mobilidade Total via USB-Serial OTG", "Funciona diretamente no celular Android comum via cabo serial OTG, sem necessidade de notebook ou tomada no rack."),
        ("Operação 100% Offline-First", "Validação completa e aplicação de scripts mesmo em locais sem sinal 4G/5G ou sem internet."),
        ("Conferência Física Prévia (Carrier Detection)", "Auditoria ativa de link que impede reloads prematuros caso o cabo físico de rede esteja desconectado."),
        ("Certidão Digital com QR Code", "Emissão na hora de laudo formal em PDF assinado digitalmente com data, hora, IPs, testes de banda e fotos da instalação.")
    ]

    for idx, (titulo, desc) in enumerate(itens_mobile):
        p = sh_card_m.text_frame.paragraphs[0] if idx == 0 else sh_card_m.text_frame.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = "• " + titulo + "\n  "
        r1.font.name = "AMX Black"
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = CLARO_RED

        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "AMX"
        r2.font.size = Pt(11)
        r2.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 9: Os 4 Pilares Estratégicos & Modelo Operacional
    # ==========================================
    print("Atualizando Slide 9 (4 Pilares Estratégicos)...")
    s9 = prs.slides[8]
    sh_t9 = [sh for sh in s9.shapes if sh.shape_id == 39][0]
    sh_t9.text_frame.paragraphs[0].text = "Os 4 Pilares Estratégicos & Modelo Operacional"
    sh_t9.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_t9.text_frame.paragraphs[0].font.size = Pt(32)
    sh_t9.text_frame.paragraphs[0].font.bold = True

    # Remover imagens e caixas antigas
    for old_id in [8, 9, 3, 11]:
        matching = [sh for sh in s9.shapes if sh.shape_id == old_id]
        if matching:
            remove_shape(s9, matching[0])

    pilares = [
        ("1. Flexibilidade de Uso (Offline-First)",
         "• Operação Híbrida: Triagem na Base (Desktop) + Validação ágil na Ponta (Mobile OTG).\n• Cache local de firmwares: ativação garantida mesmo em áreas de sombra ou sem link de internet.",
         Inches(0.85), Inches(1.5), Inches(5.6), Inches(2.4)),
        ("2. Redução Drástica de Custos (Soberania)",
         "• R$ 0 em licenciamento de software (vs custos recorrentes do Panteon).\n• Zero novos servidores: roda na infraestrutura existente de PCs e celulares.\n• Alavanca para renegociar contratos de empreiteiras (menor valor por OS).",
         Inches(6.85), Inches(1.5), Inches(5.6), Inches(2.4)),
        ("3. Agilidade & Suporte Interno",
         "• Evolução contínua pela equipe interna da Claro Regional SC.\n• Autonomia total para homologar novos modelos e versões sem depender de terceiros.\n• Resolução imediata de chamados e adaptação a demandas operacionais locais.",
         Inches(0.85), Inches(4.15), Inches(5.6), Inches(2.5)),
        ("4. Triagem Fisiológica & Visita Única",
         "• Roteador sai da base com portas, memórias e boot 100% testados.\n• Elimina 100% das visitas perdidas por roteadores com defeito físico.\n• Instalação Plug & Play no rack concluída em menos de 10 minutos.",
         Inches(6.85), Inches(4.15), Inches(5.6), Inches(2.5))
    ]

    for titulo, desc, left, top, width, height in pilares:
        # Adicionar retângulo de fundo suave
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CLARO_LIGHT_BG
        card.line.color.rgb = CLARO_BORDER
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.18)
        tf.margin_bottom = Inches(0.18)
        p = tf.paragraphs[0]
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = titulo + "\n"
        r1.font.name = "AMX Black"
        r1.font.size = Pt(13)
        r1.font.bold = True
        r1.font.color.rgb = CLARO_RED

        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "AMX"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 10: Resultados Alcançados
    # ==========================================
    print("Atualizando Slide 10 (Resultados Alcançados)...")
    s10 = prs.slides[9]
    sh_r = [sh for sh in s10.shapes if sh.shape_id == 47][0]
    sh_r.text_frame.paragraphs[0].text = "Resultados alcançados"
    sh_r.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_r.text_frame.paragraphs[0].font.size = Pt(32)
    sh_r.text_frame.paragraphs[0].font.bold = True

    sh_exp = [sh for sh in s10.shapes if sh.shape_id == 59][0]
    sh_exp.text_frame.paragraphs[0].text = "Ganhos Quantitativos & Eficiência Operacional"
    sh_exp.text_frame.paragraphs[0].font.name = "AMX Medium"
    sh_exp.text_frame.paragraphs[0].font.size = Pt(16)
    sh_exp.text_frame.paragraphs[0].font.color.rgb = CLARO_RED

    sh_left_t = [sh for sh in s10.shapes if sh.shape_id == 57][0]
    sh_left_t.width = Inches(4.2)
    sh_left_t.height = Inches(4.5)
    clear_text_frame(sh_left_t.text_frame)

    kpis_sparc = [
        ("Tempo de Ativação:", "Redução de ~50 min para <10 min por circuito (-88% de tempo)."),
        ("Aderência SAIP:", "100% de precisão, eliminação total de erros de digitação."),
        ("Visita Única (VU):", "Eliminação de retrabalho por portas queimadas ou configs residuais."),
        ("Produtividade:", "Capacidade de atendimento ampliada em até 4x por técnico."),
        ("Custo de Licenças:", "R$ 0,00 — economia direta, recorrente e sustentável para a Claro.")
    ]

    for idx, (titulo, desc) in enumerate(kpis_sparc):
        p = sh_left_t.text_frame.paragraphs[0] if idx == 0 else sh_left_t.text_frame.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + titulo + " "
        r1.font.name = "AMX Black"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = CLARO_RED

        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "AMX"
        r2.font.size = Pt(10)
        r2.font.color.rgb = CLARO_DARK

    # Métricas no grupo Agrupar 60
    grp60 = [sh for sh in s10.shapes if sh.name == 'Agrupar 60'][0]
    sh_nps = [sub for sub in grp60.shapes if sub.shape_id == 3075][0]
    sh_nps.text_frame.paragraphs[0].text = "-88%"
    sh_nps.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_nps.text_frame.paragraphs[0].font.size = Pt(28)
    sh_nps.text_frame.paragraphs[0].font.bold = True
    sh_nps.text_frame.paragraphs[0].font.color.rgb = CLARO_RED

    sh_nps_desc = [sub for sub in grp60.shapes if sub.shape_id == 3076][0]
    sh_nps_desc.text_frame.paragraphs[0].text = "Tempo de Ativação\n(de 50 min para <10 min)"
    sh_nps_desc.text_frame.paragraphs[0].font.name = "AMX"
    sh_nps_desc.text_frame.paragraphs[0].font.size = Pt(10)

    sh_lig = [sub for sub in grp60.shapes if sub.shape_id == 3077][0]
    sh_lig.text_frame.paragraphs[0].text = "100%"
    sh_lig.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_lig.text_frame.paragraphs[0].font.size = Pt(28)
    sh_lig.text_frame.paragraphs[0].font.bold = True
    sh_lig.text_frame.paragraphs[0].font.color.rgb = CLARO_RED

    sh_lig_desc = [sub for sub in grp60.shapes if sub.shape_id == 3078][0]
    sh_lig_desc.text_frame.paragraphs[0].text = "Aderência SAIP e\nZero Erro de Digitação"
    sh_lig_desc.text_frame.paragraphs[0].font.name = "AMX"
    sh_lig_desc.text_frame.paragraphs[0].font.size = Pt(10)

    # Adicionar 2 badges de métricas adicionais abaixo de Agrupar 60
    # Badge R$ 0
    b1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.02), Inches(4.90), Inches(2.3), Inches(1.8))
    b1.fill.solid()
    b1.fill.fore_color.rgb = CLARO_LIGHT_BG
    b1.line.color.rgb = CLARO_BORDER
    tf1 = b1.text_frame
    tf1.margin_top = Inches(0.15)
    p = tf1.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "R$ 0\n"
    r.font.name = "AMX Black"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = CLARO_RED
    r2 = p.add_run()
    r2.text = "Custo com Licenças\nou Novos Servidores"
    r2.font.name = "AMX"
    r2.font.size = Pt(10)
    r2.font.color.rgb = CLARO_DARK

    # Badge 0% Visitas Perdidas
    b2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.51), Inches(4.90), Inches(2.3), Inches(1.8))
    b2.fill.solid()
    b2.fill.fore_color.rgb = CLARO_LIGHT_BG
    b2.line.color.rgb = CLARO_BORDER
    tf2 = b2.text_frame
    tf2.margin_top = Inches(0.15)
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "0%\n"
    r.font.name = "AMX Black"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = CLARO_RED
    r2 = p.add_run()
    r2.text = "Visitas Perdidas por\nFalhas de Hardware"
    r2.font.name = "AMX"
    r2.font.size = Pt(10)
    r2.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 11: Impacto no Cliente
    # ==========================================
    print("Atualizando Slide 11 (Impacto no Cliente)...")
    s11 = prs.slides[10]
    sh_t11 = [sh for sh in s11.shapes if sh.shape_id == 47][0]
    sh_t11.text_frame.paragraphs[0].text = "Impacto no Cliente (Experiência B2B & Satisfação)"
    sh_t11.text_frame.paragraphs[0].font.name = "AMX Black"
    sh_t11.text_frame.paragraphs[0].font.size = Pt(32)
    sh_t11.text_frame.paragraphs[0].font.bold = True

    # Imagem à esquerda (Shape 4) -> Substituir com screenshot da Certidão Digital / Speedtest
    sh_img_c = [sh for sh in s11.shapes if sh.shape_id == 4][0]
    set_shape_image(s11, sh_img_c, IMG_MOBILE_TEST)

    # Texto à direita (Shape 3)
    sh_txt11 = [sh for sh in s11.shapes if sh.shape_id == 3][0]
    sh_txt11.left = Inches(6.6)
    sh_txt11.top = Inches(1.6)
    sh_txt11.width = Inches(6.0)
    sh_txt11.height = Inches(4.2)
    clear_text_frame(sh_txt11.text_frame)

    impactos = [
        ("A entrega do circuito é o primeiro e mais crítico contato do cliente corporativo com a Claro:", True),
        ("1. Agilidade e Respeito ao Tempo do Cliente: A ativação é concluída em <10 min no rack, liberando o circuito imediatamente sem paralisações na empresa.", False),
        ("2. Garantia Real da Visita Única (VU): O cliente nunca é surpreendido com 'equipamento defeituoso' ou 'configuração incompatível', eliminando reagendamentos.", False),
        ("3. Transparência e Confiança Imediata: O cliente recebe na hora uma Certidão Digital formal em PDF com QR Code, atestando conformidade técnica, latência e velocidade real.", False),
        ("4. Elevação Direta do NPS B2B: A primeira impressão de excelência técnica e velocidade transforma a ativação em um momento de alto encantamento.", False)
    ]

    for idx, (txt, is_header) in enumerate(impactos):
        p = sh_txt11.text_frame.paragraphs[0] if idx == 0 else sh_txt11.text_frame.add_paragraph()
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = txt
        if is_header:
            r.font.name = "AMX Black"
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = CLARO_RED
        else:
            r.font.name = "AMX"
            r.font.size = Pt(11)
            r.font.color.rgb = CLARO_DARK

    # ==========================================
    # SLIDE 12: Encerramento Institucional
    # ==========================================
    print("Atualizando Slide 12 (Encerramento)...")
    s12 = prs.slides[11]
    
    # Adicionar caixa de créditos / encerramento
    tx_end = s12.shapes.add_textbox(Inches(2.5), Inches(4.5), Inches(8.33), Inches(1.5))
    tf_end = tx_end.text_frame
    p = tf_end.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r1 = p.add_run()
    r1.text = "SPARC — Sistema de Provisionamento e Ativação de Roteadores Claro\n"
    r1.font.name = "AMX Black"
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = CLARO_WHITE

    p2 = tf_end.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "George Luís Calcmann — Analista em Serviços de Telecomunicações II\nRegional Sul — Santa Catarina | Claro Brasil"
    r2.font.name = "AMX"
    r2.font.size = Pt(13)
    r2.font.color.rgb = CLARO_WHITE

    # Salvar
    prs.save(OUT_PPT)
    print(f"Apresentação salva com sucesso em: {OUT_PPT}")

    os.makedirs(os.path.dirname(DIST_PPT), exist_ok=True)
    shutil.copy(OUT_PPT, DIST_PPT)
    print(f"Cópia distribuída para dist-beta: {DIST_PPT}")

if __name__ == '__main__':
    main()
